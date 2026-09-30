from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

# Colors
BG_COLOR = RGBColor(15, 15, 15)  # Dark background (#0F0F0F)
ACCENT_COLOR = RGBColor(249, 115, 22)  # Saffron Orange (#F97316)
TEXT_PRIMARY = RGBColor(248, 250, 252)  # Near White (#F8FAFC)
TEXT_SECONDARY = RGBColor(148, 163, 184)  # Muted Slate (#94A3B8)

def apply_theme(slide):
    # Set background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def format_title(title_shape, text):
    title_shape.text = text
    for paragraph in title_shape.text_frame.paragraphs:
        paragraph.font.color.rgb = ACCENT_COLOR
        paragraph.font.name = 'Arial'
        paragraph.font.bold = True
        paragraph.font.size = Pt(40)

def add_content(slide, left, top, width, height, text, font_size=20, color=TEXT_PRIMARY, is_bold=False):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.name = 'Arial'
    p.font.bold = is_bold
    return tf

prs = Presentation()
# Set slide dimensions to 16:9
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

# SLIDE 1: Title
slide1 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide1)
banner_path = r"d:\Hackathons\12 Samsung GenAI Hackathon\Repository\VaakSetu\assets\banner.jpg"
if os.path.exists(banner_path):
    slide1.shapes.add_picture(banner_path, Inches(0), Inches(0), width=Inches(13.333))

# Dark overlay for text legibility
overlay = slide1.shapes.add_shape(
    1, Inches(0), Inches(0), Inches(13.333), Inches(7.5) # 1 is AutoShapeType.RECTANGLE
)
overlay.fill.solid()
overlay.fill.fore_color.rgb = BG_COLOR
overlay.fill.transparency = 0.6
overlay.line.fill.background()

add_content(slide1, Inches(1), Inches(2), Inches(11), Inches(1), "VaakSetu", font_size=60, color=ACCENT_COLOR, is_bold=True)
add_content(slide1, Inches(1), Inches(3.2), Inches(11), Inches(0.5), "The Voice Bridge for Bharat", font_size=32)
add_content(slide1, Inches(1), Inches(4), Inches(11), Inches(0.5), "Interruptible Real-Time Conversational AI Platform", font_size=24, color=TEXT_SECONDARY)
add_content(slide1, Inches(1), Inches(5.5), Inches(11), Inches(1), "Team AGNISHAKTI\nShaurya Kesarwani | Sudhanshu Kumar | Md Nayaj | Mouriyan\nSRMIST Chennai | Samsung Hack2Future 2.0", font_size=18, color=TEXT_SECONDARY)

# SLIDE 2: The Problem
slide2 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide2)
add_content(slide2, Inches(1), Inches(0.5), Inches(11), Inches(1), "The Challenge of Real-Time Interruptible Voice", font_size=40, color=ACCENT_COLOR, is_bold=True)
content2 = [
    "• Standard AI operates in half-duplex (listen, think, speak)",
    "• Users frequently interrupt, correct themselves, and re-plan mid-sentence",
    "• Sequential processing causes unacceptable latency and stale conversational states",
    "• The core challenge is concurrency: running perception, reasoning, and speech on a unified timeline"
]
for i, bullet in enumerate(content2):
    add_content(slide2, Inches(1), Inches(2 + (i*0.8)), Inches(11), Inches(0.5), bullet, font_size=24)
add_content(slide2, Inches(1), Inches(5.5), Inches(11), Inches(1), '"Real conversations are not turn-based. They are concurrent."', font_size=32, color=ACCENT_COLOR, is_bold=True)

# SLIDE 3: Why Existing Tools Fail
slide3 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide3)
add_content(slide3, Inches(1), Inches(0.5), Inches(11), Inches(1), "Traditional Architecture Bottlenecks", font_size=40, color=ACCENT_COLOR, is_bold=True)

tools = [
    ("Cloud STT / TTS Platforms", "High latency in streaming mode, no built-in state rollback upon interruption"),
    ("Standard LLM Pipelines", "High Time To First Token blocks concurrent tool execution"),
    ("Traditional Dialog Trees", "Rigid turn-taking, completely fails on mid-flow self-repairs"),
    ("Generic Voice Bots", "Single execution thread, cannot drop stale routing paths")
]
top_margin = 2.0
for i, (tool, detail) in enumerate(tools):
    add_content(slide3, Inches(1), Inches(top_margin + (i*1.1)), Inches(4), Inches(1), tool, font_size=24, is_bold=True)
    add_content(slide3, Inches(5), Inches(top_margin + (i*1.1)), Inches(7), Inches(1), detail, font_size=22, color=TEXT_SECONDARY)

# SLIDE 4: Our Solution
slide4 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide4)
add_content(slide4, Inches(1), Inches(0.5), Inches(11), Inches(1), "Dual-Process Concurrency for Voice", font_size=40, color=ACCENT_COLOR, is_bold=True)

cols = [
    ("FAST PATH", "Responsiveness\n\nLow-latency acknowledgments and conversational fillers", Inches(1)),
    ("SLOW PATH", "Asynchronous Reasoning\n\nComplex reasoning, tool execution, and data extraction", Inches(5)),
    ("COORDINATION LAYER", "State Management\n\nHandles non-blocking execution and call cancellation", Inches(9))
]
for title, desc, left in cols:
    add_content(slide4, left, Inches(2), Inches(3.5), Inches(0.5), title, font_size=24, color=ACCENT_COLOR, is_bold=True)
    add_content(slide4, left, Inches(2.8), Inches(3.5), Inches(2), desc, font_size=20)

differentiators = [
    "• Parallel Execution Architecture: Fast Path and Slow Path run concurrently",
    "• Dynamic State Snapshots: Emits structured JSON state dynamically",
    "• Real-Time Interruption Handling: Cancels stale in-flight reasoning",
    "• RLAIF Self-Improvement: Automated evaluation of agent performance"
]
for i, diff in enumerate(differentiators):
    add_content(slide4, Inches(1), Inches(4.5 + (i*0.6)), Inches(11), Inches(0.5), diff, font_size=20, color=TEXT_SECONDARY)

# SLIDE 5: Architecture
slide5 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide5)
add_content(slide5, Inches(1), Inches(0.5), Inches(11), Inches(1), "End-to-End System Architecture", font_size=40, color=ACCENT_COLOR, is_bold=True)
add_content(slide5, Inches(1), Inches(1.8), Inches(11), Inches(3), 
    "• WebRTC / WebSocket ingress for continuous streaming\n"
    "• Transcription and Language ID layer\n"
    "• Parallel Orchestration via LangGraph\n"
    "• State Snapshot persistence to Database\n\n"
    "[Insert Image: End-to-End Architecture Diagram detailing Fast Path, Slow Path, and Coordination Layer]", 
    font_size=24)
add_content(slide5, Inches(1), Inches(6), Inches(11), Inches(1), 
    "Python 3.10 | FastAPI | Next.js 15 | React 19 | Firebase | Sarvam AI | Gemini 2.0 Flash | LangGraph", 
    font_size=20, color=ACCENT_COLOR, is_bold=True)

# SLIDE 6: Fast-and-Slow Execution Pipeline
slide6 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide6)
add_content(slide6, Inches(1), Inches(0.5), Inches(11), Inches(1), "Concurrency in Action", font_size=40, color=ACCENT_COLOR, is_bold=True)
add_content(slide6, Inches(1), Inches(1.8), Inches(11), Inches(4), 
    "Fast Path: Floor Management\n"
    "Issues meaningful responses quickly without making false completion claims. Manages audio playback and conversational fillers.\n\n"
    "Slow Path: Structured Extraction\n"
    "Parses dynamic tool definitions. Extracts data using LLMs without blocking the fast path.\n\n"
    "Coordination: Interruption Recovery\n"
    "Cancels superseded in-flight tool calls. Updates state snapshots and re-plans cleanly upon user interruption.\n\n"
    "[Insert Image: Working Pipeline showing Concurrent Execution and Interruption Signals]", 
    font_size=22)
add_content(slide6, Inches(1), Inches(6.2), Inches(11), Inches(0.5), "Parallel execution reduces perceived latency drastically while maintaining state consistency.", font_size=22, color=ACCENT_COLOR, is_bold=True)

# SLIDE 7: Dynamic State Snapshots
slide7 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide7)
add_content(slide7, Inches(1), Inches(0.5), Inches(11), Inches(1), "Structured State Tracking in Real-Time", font_size=40, color=ACCENT_COLOR, is_bold=True)
add_content(slide7, Inches(1), Inches(2), Inches(11), Inches(3), 
    "• Session-Scoped Memory: Maintains intent and slot values across multi-turn interactions\n"
    "• Incremental Extraction: Maps user speech to specific schemas (e.g., Healthcare, Finance) concurrently\n"
    "• Protocol Compliance: Emits well-formed JSON payloads (State Snapshots) upon turn completion\n"
    "• Idempotent Updates: Strictly avoids duplicate state-changing calls during rapid interruptions\n\n"
    "[Insert Image: Pipeline diagram showing raw speech turning into JSON State Snapshots]", 
    font_size=24)

# SLIDE 8: Automated Quality Evaluation (RLAIF)
slide8 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide8)
add_content(slide8, Inches(1), Inches(0.5), Inches(11), Inches(1), "Evaluating Interruptible Agents", font_size=40, color=ACCENT_COLOR, is_bold=True)

add_content(slide8, Inches(1), Inches(2), Inches(5.5), Inches(0.5), "40% Programmatic (Instant)", font_size=26, color=ACCENT_COLOR, is_bold=True)
add_content(slide8, Inches(1), Inches(2.8), Inches(5.5), Inches(2), "• Response Latency\n• Interruption Recovery Speed\n• Cancellation of Invalidated Calls\n• Absence of Stale Re-runs", font_size=24)

add_content(slide8, Inches(7), Inches(2), Inches(5.5), Inches(0.5), "60% LLM Judge (Gemini)", font_size=26, color=ACCENT_COLOR, is_bold=True)
add_content(slide8, Inches(7), Inches(2.8), Inches(5.5), Inches(2), "• Task Completion Accuracy\n• State Snapshot Accuracy\n• Protocol Compliance\n• Final Response Grounding", font_size=24)

add_content(slide8, Inches(1), Inches(5.5), Inches(11), Inches(1), "Output:\nScore above 0.70: DPO label chosen\nScore below 0.70: DPO label rejected", font_size=24, color=TEXT_SECONDARY, is_bold=True)

# SLIDE 9: Applied Real-Time Agents
slide9 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide9)
add_content(slide9, Inches(1), Inches(0.5), Inches(11), Inches(1), "Schema-Driven Domain Switching", font_size=40, color=ACCENT_COLOR, is_bold=True)

add_content(slide9, Inches(1), Inches(2), Inches(5.5), Inches(0.5), "Healthcare", font_size=28, color=ACCENT_COLOR, is_bold=True)
add_content(slide9, Inches(1), Inches(2.8), Inches(5.5), Inches(2), "• Extracting symptoms, diagnosis, and treatment plan mid-consultation\n• Resolving speech hesitations and conversational self-repairs\n• Gracefully handling interruptions without dropping prior context", font_size=22)

add_content(slide9, Inches(7), Inches(2), Inches(5.5), Inches(0.5), "Financial Services", font_size=28, color=ACCENT_COLOR, is_bold=True)
add_content(slide9, Inches(7), Inches(2.8), Inches(5.5), Inches(2), "• Adjusting parameters mid-interaction (e.g., changing payment promises)\n• Strictly avoiding duplicate state-modifying calls (double-booking)\n• Emitting updated State Snapshots instantly", font_size=22)

add_content(slide9, Inches(1), Inches(6.2), Inches(11), Inches(0.5), "Schema-Driven Tools: Switch domains with one config change. No retraining required.", font_size=24, color=ACCENT_COLOR, is_bold=True)

# SLIDE 10: Live Demo Flow
slide10 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide10)
add_content(slide10, Inches(1), Inches(0.5), Inches(11), Inches(1), "Execution & Interruption Handling Demo", font_size=40, color=ACCENT_COLOR, is_bold=True)
flow = [
    "1. Open dashboard, select domain",
    "2. Start streaming audio via WebSocket",
    "3. User speaks, pauses, and interrupts themselves",
    "4. Agent executes Fast Path acknowledgment",
    "5. Agent updates State Snapshot via Slow Path extraction",
    "6. Interruption triggers cancellation of stale API calls",
    "7. Final JSON payload emitted with correct intent and slots",
    "8. View RLAIF evaluation score post-session"
]
for i, step in enumerate(flow):
    add_content(slide10, Inches(1), Inches(2 + (i*0.45)), Inches(11), Inches(0.4), step, font_size=20)
add_content(slide10, Inches(1), Inches(6.0), Inches(11), Inches(0.5), "[Insert Image: Working Pipeline demonstrating a mid-sentence interruption recovery]", font_size=20, color=TEXT_SECONDARY)

# SLIDE 11: Frontend: The Control Dashboard
slide11 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide11)
add_content(slide11, Inches(1), Inches(0.5), Inches(11), Inches(1), "Monitoring Concurrent Execution", font_size=40, color=ACCENT_COLOR, is_bold=True)

add_content(slide11, Inches(1), Inches(1.8), Inches(11), Inches(4), 
    "Agent Builder\n"
    "Define schemas and parameters for Slow Path extraction. Configures the state payload structure.\n\n"
    "Session Interface\n"
    "View Fast Path responses alongside live State Snapshots updating in real-time.\n\n"
    "Trace Logging & Analytics\n"
    "Monitor latency to first substantive action, transcription chunks, and tool execution status.\n\n"
    "[Insert Image: Dashboard UI showing live transcription, latency metrics, and state extraction]", 
    font_size=22)

# SLIDE 12: Business Impact & Scale
slide12 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide12)
add_content(slide12, Inches(1), Inches(0.5), Inches(11), Inches(1), "Production-Ready Interruptible Voice AI", font_size=40, color=ACCENT_COLOR, is_bold=True)

add_content(slide12, Inches(1), Inches(2), Inches(11), Inches(2), 
    "• Drastically reduces Time-To-First-Action for complex voice queries\n"
    "• Zero duplicate state-changing calls during rapid user interruptions\n"
    "• Seamlessly handles code-mixed Indian languages\n"
    "• Built on open research: RLAIF (Google Brain), DPO (Stanford)", font_size=24)

add_content(slide12, Inches(1), Inches(4.2), Inches(11), Inches(1.5), 
    "Team AGNISHAKTI:\n"
    "• Shaurya Kesarwani | AI Core (Fast/Slow Execution)\n"
    "• Sudhanshu Kumar | Backend (WebSockets & Routes)\n"
    "• Md Nayaj | Frontend (Dashboard & Real-Time UI)\n"
    "• Mouriyan | Output Layer (TTS & Simulation)", font_size=20, color=TEXT_SECONDARY)

add_content(slide12, Inches(1), Inches(6.5), Inches(11), Inches(0.5), "VaakSetu — Bridging Languages, Mastering Real-Time Concurrency.", font_size=28, color=ACCENT_COLOR, is_bold=True)

output_path = r"d:\Hackathons\12 Samsung GenAI Hackathon\Repository\VaakSetu\VaakSetu_Samsung_Hackathon.pptx"
prs.save(output_path)
print(f"PPT saved to {output_path}")
